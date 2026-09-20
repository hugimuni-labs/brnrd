#!/usr/bin/env python3
"""Build the brnrd investor deck v1 from spec.md (never hand-edit XML).

Look = the v11 build (Aptos, ink/amber/cool palette, dimmed full-bleed stills).
Assets are read from ./assets. Resident commit share is measured live via git.
"""
from __future__ import annotations

import argparse
import os
import re
import shutil
import subprocess
import tempfile
import zipfile
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt

HERE = Path(__file__).resolve().parent
A = HERE / "assets"
FONT = "Aptos"
INK = RGBColor(0xF6, 0xF8, 0xFC)
MUTED = RGBColor(0xBE, 0xC6, 0xD2)
COOL = RGBColor(0xC6, 0xE5, 0xFF)
AMBER = RGBColor(0xFF, 0xB7, 0x39)
DIM = RGBColor(0x88, 0x92, 0xA2)
GROUND = RGBColor(0x08, 0x0A, 0x0F)
RED = RGBColor(0xFF, 0x8A, 0x7A)
FIXED_TS = (2026, 9, 20, 0, 0, 0)


# ---------------------------------------------------------------- spec
def git(*args: str) -> str:
    return subprocess.run(["git", *args], check=True, capture_output=True, text=True).stdout.strip()


def measure() -> dict[str, str]:
    bot = int(git("log", "--author=brnrd-bot", "--oneline", "main").count("\n") + 1)
    total = int(git("rev-list", "--count", "main"))
    return {"bot": f"{bot:,}", "total": f"{total:,}", "pct": f"{round(100 * bot / total)} %"}


def parse_spec(path: Path, subs: dict[str, str]) -> dict[str, dict]:
    text = path.read_text(encoding="utf-8")
    for k, v in subs.items():
        text = text.replace("{" + k + "}", v)
    slides: dict[str, dict] = {}
    cur = None
    for line in text.splitlines():
        m = re.match(r"## (\S+)$", line)
        if m:
            cur = slides[m.group(1)] = {"label": "", "headline": "", "body": [], "notes": []}
        elif cur is not None:
            if line.startswith("label: "):
                cur["label"] = line[7:]
            elif line.startswith("headline: "):
                cur["headline"] = line[10:]
            elif line.startswith("- "):
                cur["body"].append([p.strip() for p in line[2:].split("|")])
            elif line.startswith("> "):
                cur["notes"].append(line[2:])
    return slides


# ---------------------------------------------------------------- drawing
def pic(slide, path, l, t, w, h):
    return slide.shapes.add_picture(str(path), Inches(l), Inches(t), Inches(w), Inches(h))


def rect(slide, l, t, w, h, color, opacity=1.0):
    s = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(l), Inches(t), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = color
    s.line.fill.background()
    if opacity < 1.0:
        a = OxmlElement("a:alpha")
        a.set("val", str(round(opacity * 100000)))
        s._element.spPr.solidFill.srgbClr.append(a)
    return s


def text(slide, s, l, t, w, h, *, size, color=INK, bold=False, italic=False, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, spacing=None):
    tb = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    f = tb.text_frame
    f.margin_left = f.margin_right = f.margin_top = f.margin_bottom = 0
    f.vertical_anchor = anchor
    f.word_wrap = True
    p = f.paragraphs[0]
    p.alignment = align
    if spacing:
        p.line_spacing = spacing
    r = p.add_run()
    r.text = s
    r.font.name = FONT
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.italic = italic
    r.font.color.rgb = color
    return tb


def bg(slide, name, opacity):
    pic(slide, A / name, 0, 0, 13.333333, 7.5)
    rect(slide, 0, 0, 13.333333, 7.5, GROUND, opacity)


def chrome(slide, label, number):
    text(slide, label, 0.68, 0.55, 9.0, 0.25, size=11.5, color=COOL, bold=True)
    text(slide, "brnrd >_", 0.55, 6.96, 1.3, 0.18, size=10.5)
    rect(slide, 2.25, 7.075, 0.42, 0.015, AMBER)
    text(slide, "hugimuni sas · investor deck v1", 9.5, 6.98, 2.7, 0.16, size=8.5, color=DIM, align=PP_ALIGN.RIGHT)
    text(slide, number, 12.5, 6.98, 0.4, 0.16, size=8.5, color=DIM, align=PP_ALIGN.RIGHT)


def headline(slide, s, width=11.9, size=30, top=0.98, height=1.2):
    text(slide, s, 0.68, top, width, height, size=size, spacing=0.95)


def movie(slide, name, tmp, rect_):
    poster = tmp / f"{name}.png"
    subprocess.run(["ffmpeg", "-hide_banner", "-loglevel", "error", "-y", "-i", str(A / name),
                    "-vf", "select=eq(n\\,30)", "-frames:v", "1", str(poster)], check=True)
    l, t, w, h = rect_
    return slide.shapes.add_movie(str(A / name), Inches(l), Inches(t), Inches(w), Inches(h),
                                  poster_frame_image=str(poster), mime_type="video/mp4")


def bar_rows(slide, rows, l, t, w, gap, *, size=15, title_color=INK, sub=True, colors=(COOL, AMBER)):
    y = t
    for i, row in enumerate(rows):
        c = colors[i % 2]
        rect(slide, l, y, 0.045, 0.62 if sub and len(row) > 1 else 0.4, c)
        text(slide, row[0], l + 0.22, y - 0.02, w, 0.36, size=size, color=title_color, bold=True)
        if sub and len(row) > 1:
            text(slide, row[1], l + 0.22, y + 0.3, w, 0.5, size=size - 3, color=MUTED)
        y += gap


def stat(slide, big, small, l, t, w, color=AMBER, bigsize=30):
    text(slide, big, l, t, w, 0.6, size=bigsize, color=color, bold=True)
    text(slide, small, l, t + 0.62, w, 0.6, size=12, color=MUTED)


# ---------------------------------------------------------------- slides
def build(prs, S, tmp):
    blank = prs.slide_layouts[6]
    new = lambda: prs.slides.add_slide(blank)

    def notes(slide, key):
        slide.notes_slide.notes_text_frame.text = "\n".join(S[key]["notes"])

    # 1 sentence
    d = S["1"]; s = new(); bg(s, "bg-dark.jpg", 0.60); chrome(s, d["label"], "01")
    headline(s, d["headline"], width=7.6, size=38, top=1.35, height=1.7)
    rect(s, 0.72, 3.25, 2.4, 0.018, AMBER)
    text(s, d["body"][0][0], 0.72, 3.55, 7.0, 1.8, size=19, color=MUTED, spacing=1.08)
    text(s, d["body"][1][0], 0.72, 5.5, 6.0, 0.3, size=14, color=AMBER, bold=True)
    movie(s, "loop-3-stage.mp4", tmp, (8.35, 1.2, 4.5, 2.53))
    text(s, d["body"][2][0], 8.35, 3.9, 4.5, 0.3, size=11.5, color=COOL, italic=True)
    notes(s, "1")

    # 2 proof
    d = S["2"]; s = new(); bg(s, "bg-blue.jpg", 0.84); chrome(s, d["label"], "02")
    headline(s, d["headline"], size=32, height=0.7)
    cells = d["body"][:6]
    for i, (big, small) in enumerate(cells):
        col, row = i % 3, i // 3
        stat(s, big, small, 0.75 + col * 4.05, 2.1 + row * 1.75, 3.7, color=AMBER if i in (3,) else INK, bigsize=32)
        rect(s, 0.75 + col * 4.05, 2.02 + row * 1.75, 0.6, 0.02, AMBER)
    text(s, d["body"][6][0], 0.75, 5.75, 11.5, 0.5, size=20, color=AMBER, italic=True)
    notes(s, "2")

    # 3 why now
    d = S["3"]; s = new(); bg(s, "bg-wide.jpg", 0.84); chrome(s, d["label"], "03")
    headline(s, d["headline"], width=11.6, size=27, height=1.0)
    text(s, "Their own docs", 0.75, 2.05, 5, 0.3, size=12, color=COOL, bold=True)
    bar_rows(s, d["body"][:4], 0.75, 2.5, 5.6, 0.95, size=14.5)
    text(s, "Their own trackers", 7.0, 2.05, 5, 0.3, size=12, color=COOL, bold=True)
    for i, (big, small) in enumerate(d["body"][4:7]):
        stat(s, big, small, 7.0, 2.5 + i * 1.12, 5.6, bigsize=26)
    text(s, d["body"][7][0], 0.75, 6.2, 8.0, 0.4, size=17, color=AMBER, italic=True)
    text(s, d["body"][8][0], 8.6, 6.3, 4.0, 0.3, size=12, color=MUTED, align=PP_ALIGN.RIGHT)
    notes(s, "3")

    # 4 wedge
    d = S["4"]; s = new(); bg(s, "bg-dark.jpg", 0.83); chrome(s, d["label"], "04")
    headline(s, d["headline"], width=7.0, size=30, height=0.7)
    bar_rows(s, d["body"], 0.75, 1.95, 7.0, 0.98, size=15.5)
    movie(s, "loop-2-the-wall.mp4", tmp, (8.35, 1.95, 4.5, 2.8))
    text(s, "quota wall, 5 Sep 18:16 — core flipped, same thread", 8.35, 4.9, 4.5, 0.3, size=11.5, color=COOL, italic=True)
    notes(s, "4")

    # 5 product
    d = S["5"]; s = new(); bg(s, "bg-blue.jpg", 0.84); chrome(s, d["label"], "05")
    headline(s, d["headline"], width=7.2, size=30, height=1.3)
    bar_rows(s, d["body"], 0.75, 2.75, 6.8, 0.98, size=17)
    movie(s, "loop-4-away.mp4", tmp, (8.1, 1.15, 4.7, 3.02))
    text(s, "the seat, overnight — run cards, not a log", 8.1, 4.35, 4.7, 0.3, size=11.5, color=COOL, italic=True)
    notes(s, "5")

    # 6 place
    d = S["6"]; s = new(); bg(s, "bg-dark.jpg", 0.85); chrome(s, d["label"], "06")
    headline(s, d["headline"], width=5.6, size=32, height=1.3)
    text(s, d["body"][0][0], 0.75, 2.45, 4.9, 1.6, size=15.5, color=MUTED, spacing=1.08)
    text(s, d["body"][1][0], 0.75, 4.1, 4.9, 1.5, size=16.5, color=INK, spacing=1.08)
    text(s, d["body"][2][0].upper(), 0.75, 5.95, 4.9, 0.3, size=12, color=AMBER, bold=True)
    pic(s, A / "field8-lit.png", 6.0, 0.95, 6.9, 4.3125)
    text(s, "field 8 — the streets: the tree as a place, one frame", 6.0, 5.4, 6.9, 0.3, size=11.5, color=COOL, italic=True)
    notes(s, "6")

    # 7 business model
    d = S["7"]; s = new(); bg(s, "bg-wide.jpg", 0.84); chrome(s, d["label"], "07")
    headline(s, d["headline"], size=34, height=0.7)
    rows = d["body"][:4]
    y = 1.95
    for i, r in enumerate(rows):
        c = (COOL, AMBER)[i % 2]
        if len(r) == 2:
            rect(s, 0.75, y, 0.045, 0.85, c)
            text(s, r[0], 1.0, y - 0.02, 1.6, 0.35, size=15, color=c, bold=True)
            text(s, r[1], 2.7, y - 0.02, 9.3, 0.9, size=15, color=INK, spacing=1.05)
            y += 1.05
        else:
            rect(s, 0.75, y, 0.045, 0.4, c)
            text(s, r[0], 1.0, y - 0.02, 11, 0.4, size=15, color=INK)
            y += 0.7
    text(s, d["body"][4][0], 0.75, 6.2, 11.5, 0.4, size=17, color=AMBER, italic=True)
    notes(s, "7")

    # 8 traction
    d = S["8"]; s = new(); bg(s, "bg-blue.jpg", 0.85); chrome(s, d["label"], "08")
    headline(s, d["headline"], size=40, height=0.8)
    for i, (big, small) in enumerate(d["body"][:4]):
        stat(s, big, small, 0.75 + i * 3.1, 2.15, 2.9, color=AMBER, bigsize=30)
    rect(s, 0.75, 4.05, 11.6, 0.018, AMBER)
    text(s, d["body"][4][0], 0.75, 4.3, 4.0, 0.35, size=13, color=COOL, bold=True)
    text(s, d["body"][4][1], 0.75, 4.75, 11.4, 0.5, size=19, color=INK)
    text(s, d["body"][5][0], 0.75, 5.65, 11.4, 0.5, size=20, color=AMBER, italic=True)
    notes(s, "8")

    # 9 GTM
    d = S["9"]; s = new(); bg(s, "bg-dark.jpg", 0.84); chrome(s, d["label"], "09")
    headline(s, d["headline"], size=36, height=0.7)
    text(s, "channel", 0.98, 1.95, 4, 0.3, size=11, color=COOL, bold=True)
    text(s, "its row in the funnel", 8.6, 1.95, 4, 0.3, size=11, color=COOL, bold=True)
    for i, (a, b, c) in enumerate(d["body"]):
        y = 2.4 + i * 1.02
        rect(s, 0.75, y, 0.045, 0.62, (COOL, AMBER)[i % 2])
        text(s, a, 0.98, y - 0.02, 7.3, 0.36, size=16, bold=True)
        text(s, b, 0.98, y + 0.34, 7.3, 0.3, size=12.5, color=MUTED)
        text(s, c, 8.6, y + 0.02, 4.2, 0.4, size=15, color=AMBER)
    notes(s, "9")

    # 10 competition
    d = S["10"]; s = new(); bg(s, "bg-wide.jpg", 0.85); chrome(s, d["label"], "10")
    headline(s, d["headline"], size=30, height=0.7)
    cols = [(0.75, AMBER), (4.85, COOL), (8.95, RED)]
    for (x, c), row in zip(cols, d["body"]):
        rect(s, x, 1.95, 3.6, 0.03, c)
        text(s, row[0], x, 2.12, 3.6, 0.4, size=16, color=c, bold=True)
        for j, item in enumerate(row[1:]):
            text(s, item, x, 2.75 + j * 0.85, 3.6, 0.8, size=14.5, color=INK, spacing=1.0)
    notes(s, "10")

    # 11 team
    d = S["11"]; s = new(); bg(s, "bg-wide.jpg", 0.84); chrome(s, d["label"], "11")
    headline(s, d["headline"], width=8.7, size=32, height=1.2)
    pic(s, A / "v10-mark.png", 11.4, 0.95, 0.952, 0.593)
    for i, (name, role) in enumerate(d["body"][:4]):
        x = 0.75 + i * 3.05
        c = AMBER if i == 2 else COOL
        rect(s, x, 2.75, 2.75, 0.03, c)
        text(s, name, x, 3.0, 2.75, 0.4, size=17, color=INK, bold=True)
        text(s, role, x, 3.5, 2.75, 1.5, size=14, color=AMBER if i == 2 else MUTED, spacing=1.05)
    text(s, d["body"][4][0], 0.75, 5.6, 11.5, 0.4, size=17, color=AMBER, italic=True)
    notes(s, "11")

    # 12 ask
    d = S["12"]; s = new(); bg(s, "bg-dark.jpg", 0.85); chrome(s, d["label"], "12")
    headline(s, d["headline"], size=32, height=0.7)
    for i, (k, t1, t2) in enumerate(d["body"][:3]):
        y = 1.9 + i * 1.05
        c = (AMBER, COOL, MUTED)[i]
        rect(s, 0.75, y, 0.045, 0.75, c)
        text(s, k, 1.0, y, 0.6, 0.4, size=16, color=c, bold=True)
        text(s, t1, 1.6, y - 0.02, 5.0, 0.7, size=18, color=INK, bold=True)
        text(s, t2, 6.7, y, 6.0, 0.9, size=14, color=MUTED, spacing=1.05)
    rect(s, 0.75, 5.2, 11.9, 0.018, AMBER)
    for i, (k, v) in enumerate(d["body"][3:5]):
        x = 0.75 + i * 6.1
        text(s, k.upper(), x, 5.4, 5, 0.3, size=12, color=AMBER, bold=True)
        text(s, v, x, 5.75, 5.7, 0.9, size=14, color=INK, spacing=1.05)
    notes(s, "12")

    # backups
    for key, lab in (("B1", "B1"), ("B2", "B2")):
        d = S[key]; s = new(); bg(s, "bg-dark.jpg", 0.82); chrome(s, d["label"], lab)
        headline(s, d["headline"], width=8.4, size=30, height=1.4, top=1.08)
        text(s, d["body"][0][0], 0.72, 2.75, 8.2, 2.5, size=19, color=MUTED, spacing=1.1)
        notes(s, key)


# ---------------------------------------------------------------- output
def normalize_zip(out: Path):
    tmp = out.with_suffix(".n.pptx")
    with zipfile.ZipFile(out) as src, zipfile.ZipFile(tmp, "w") as dst:
        for i in src.infolist():
            n = zipfile.ZipInfo(i.filename, FIXED_TS)
            n.compress_type = i.compress_type
            n.external_attr = i.external_attr
            dst.writestr(n, src.read(i.filename))
    os.replace(tmp, out)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--out", type=Path, default=HERE / "brnrd_investor_deck_v1.pptx")
    ap.add_argument("--desktop", type=Path, default=Path.home() / "Desktop" / "brnrd_investor_deck_v1.pptx")
    a = ap.parse_args()
    subs = measure()
    print("measured:", subs)
    S = parse_spec(HERE / "spec.md", subs)
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333333), Inches(7.5)
    prs.core_properties.title = "brnrd — investor deck v1"
    prs.core_properties.author = "HugiMuni SAS"
    with tempfile.TemporaryDirectory() as t:
        build(prs, S, Path(t))
        prs.save(a.out)
    normalize_zip(a.out)
    assert len(Presentation(a.out).slides) == 14
    if a.desktop:
        shutil.copyfile(a.out, a.desktop)
    print("built", a.out, a.out.stat().st_size)


if __name__ == "__main__":
    main()
